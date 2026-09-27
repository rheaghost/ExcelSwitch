import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from pathlib import Path


def run(task):

    input_file = task["input"]
    output_file = task["output"]

    df = pd.read_excel(input_file)

    required_columns = [
        "Region",
        "Price",
        "Quantity"
    ]

    for column in required_columns:
        if column not in df.columns:
            raise ValueError(
                f"Required column '{column}' does not exist."
            )

    # Calculate total sales per row.
    df["TotalSales"] = df["Price"] * df["Quantity"]

    total_sales = df["TotalSales"].sum()
    average_sales = df["TotalSales"].mean()

    sales_by_region = (
        df.groupby("Region")["TotalSales"]
        .sum()
        .reset_index()
    )

    # Create workbook with original data and summary.
    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            sheet_name="Data",
            index=False
        )

        sales_by_region.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
            startrow=4
        )

    # Open workbook for formatting/chart creation.
    workbook = load_workbook(output_file)

    summary = workbook["Summary"]

    summary["A1"] = "Sales Report"
    summary["A1"].font = Font(
        bold=True,
        size=16
    )

    summary["A2"] = "Total Sales"
    summary["B2"] = total_sales

    summary["A3"] = "Average Sale"
    summary["B3"] = average_sales

    # Create bar chart.
    chart = BarChart()

    chart.title = "Sales by Region"
    chart.y_axis.title = "Sales"
    chart.x_axis.title = "Region"

    data = Reference(
        summary,
        min_col=2,
        min_row=5,
        max_row=4 + len(sales_by_region)
    )

    categories = Reference(
        summary,
        min_col=1,
        min_row=6,
        max_row=4 + len(sales_by_region)
    )

    chart.add_data(
        data,
        titles_from_data=True
    )

    chart.set_categories(categories)

    summary.add_chart(
        chart,
        "D2"
    )

    workbook.save(output_file)

    print("\nReport created.")
    print(f"Rows        : {len(df)}")
    print(f"Total sales : {total_sales:.2f}")
    print(f"Average sale: {average_sales:.2f}")
    print(f"Output      : {output_file}")

    return {
        "success": True,
        "task_id": "029",
        "input": input_file,
        "output": output_file,
        "rows": len(df),
        "total_sales": float(total_sales),
        "average_sales": float(average_sales)
    }