from pathlib import Path

from openpyxl import load_workbook, Workbook


def run(task):

    input_file = task["input"]
    output_file = task["output"]

    input_path = Path(input_file)
    output_path = Path(output_file)

    if not input_path.is_file():
        raise FileNotFoundError(
            f"Input file does not exist: {input_file}"
        )

    # Read XLSX in streaming/read-only mode.
    source_wb = load_workbook(
        input_file,
        read_only=True,
        data_only=False
    )

    # Write XLSX in streaming/write-only mode.
    output_wb = Workbook(write_only=True)

    sheets_processed = 0
    rows_processed = 0

    try:

        for source_ws in source_wb.worksheets:

            target_ws = output_wb.create_sheet(
                title=source_ws.title
            )

            sheet_rows = 0

            for row in source_ws.iter_rows(
                values_only=True
            ):

                target_ws.append(list(row))

                rows_processed += 1
                sheet_rows += 1

            print(
                f"Processed sheet '{source_ws.title}': "
                f"{sheet_rows} rows"
            )

            sheets_processed += 1

        output_wb.save(output_path)

    finally:
        source_wb.close()

    print("\nLarge Excel processing complete.")
    print(f"Sheets processed: {sheets_processed}")
    print(f"Rows processed  : {rows_processed}")
    print(f"Output          : {output_file}")

    return {
        "success": True,
        "task_id": "025",
        "input": input_file,
        "output": str(output_path),
        "sheets_processed": sheets_processed,
        "rows_processed": rows_processed
    }