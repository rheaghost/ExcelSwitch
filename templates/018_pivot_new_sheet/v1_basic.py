import pandas as pd
from openpyxl import load_workbook


def run(task):

    input_file = task["input"]
    output_file = task["output"]

    index_column = task["index_column"]
    columns_column = task["columns_column"]
    values_column = task["values_column"]

    sheet_name = task.get(
        "sheet_name",
        "Pivot Summary"
    )

    source_sheet = task.get(
        "source_sheet",
        0
    )

    # Read the original worksheet
    df = pd.read_excel(
        input_file,
        sheet_name=source_sheet
    )

    required_columns = [
        index_column,
        columns_column,
        values_column
    ]

    for column in required_columns:
        if column not in df.columns:
            raise ValueError(
                f"Column '{column}' does not exist."
            )

    # Create pivot summary
    pivot = pd.pivot_table(
        df,
        index=index_column,
        columns=columns_column,
        values=values_column,
        aggfunc="sum",
        fill_value=0
    )

    pivot = pivot.reset_index()

    # Copy original workbook to output while
    # preserving its existing worksheets.
    workbook = load_workbook(input_file)

    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    workbook.save(output_file)

    # Add the pivot summary as a new worksheet
    with pd.ExcelWriter(
        output_file,
        engine="openpyxl",
        mode="a"
    ) as writer:

        pivot.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )

    print("\nPivot summary added.")
    print(f"Source sheet   : {source_sheet}")
    print(f"Index column   : {index_column}")
    print(f"Columns column : {columns_column}")
    print(f"Values column  : {values_column}")
    print(f"New sheet      : {sheet_name}")
    print(f"Output         : {output_file}")

    return {
        "success": True,
        "task_id": "018",
        "input": input_file,
        "output": output_file,
        "source_sheet": source_sheet,
        "index_column": index_column,
        "columns_column": columns_column,
        "values_column": values_column,
        "sheet_name": sheet_name
    }